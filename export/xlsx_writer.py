"""版面模型 → A3 橫式 .xlsx（openpyxl）。

與 ``pdf_writer`` 吃同一份 :class:`lib.layout_model.Sheet`，兩邊才會長一樣。

⚠️ **X 軸是人名（欄），Y 軸是日期（列）**。第一版做反了，見 layout_model 的說明。

⚠️ 這裡只負責「把版面模型畫出來」，不做任何排班判斷。要改番號怎麼算請去
``lib/rota.py``；要改表怎麼排請去 ``lib/layout_model.py``。
"""
from __future__ import annotations

from dataclasses import dataclass
from itertools import groupby

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from lib.layout_model import (
    BLUE,
    COL_BLANK,
    COL_MEMBER,
    COL_TITLE,
    RED,
    Block,
    Sheet,
    column_weight,
    GAP_CHAR,
    GAP_SCALE,
    header_spans_code_row,
    vertical_pieces,
)

FONT_NAME = "標楷體"
# 最左標題欄（維護者 2026-09-17：要看得出是標題）。中文與數字分開定字級：中文佔欄寬
# XLSX_FILL、上限 TITLE_MAX_SIZE；數字另外縮到塞得進欄寬（_digit_size），不拖累中文。
# 整串排不下欄高時中文一起縮（行高照 VERTICAL_LINE_RATIO 估）。
TITLE_MAX_SIZE = 32
# 純半形字串（番號、日期、代碼、標題的 115）用 Tahoma，與 pdf_writer.DIGIT_FAMILY 一致（維護者選 Tahoma：好認、比 Verdana 省寬度）
DIGIT_FONT_NAME = "Tahoma"


def _font_name(text) -> str:
    text = text or ""
    return DIGIT_FONT_NAME if text.isascii() and text.strip() else FONT_NAME

# ⚠️ **格子裡的字要盡量大**（維護者 2026-09-16：老人家眼睛不好，佔滿 80～90%，
# 以不切字為主；Excel 與 PDF 可以分開設定）。舊寫法一律 12pt，不管格子多大。
#
# openpyxl 量不到字寬，只能依格子點數估：標楷體中文一個字佔一個字級寬、半形
# 數字字母佔半個；橫書一行的高度約字級 × LINE_RATIO。字級取寬、高兩個限制的
# 較小者 × XLSX_FILL。估不準的部分由 Excel 的「縮小字型以適合欄寬」兜底，
# 所以不會切字，最壞是某幾格縮小一點。
# 比例要調：上機看完直接改 XLSX_FILL（PDF 另有 pdf_writer.FILL_RATIO）。
XLSX_FILL = 0.85
LINE_RATIO = 1.0
MAX_FONT_SIZE = 36
# 姓名列高度以幾個字的姓名為準。比這長的名字（複姓、原住民姓名）只縮那一格，
# 不把整列撐高——撐高會把每天的格子壓扁，整張表的字都跟著變小。
NAME_ROW_CHARS = 4

# ⚠️ 版面尺寸是**算出來的，不是猜的**——目標是自然尺寸剛好貼近 A3 橫式的
# 可列印區，讓 fitToPage 幾乎不用縮。
#
# 第一版沒算，欄寬加總超過頁寬、列高加總只有頁高的七成，結果是：Excel 依
# 寬度把整張表縮小（fitToPage **只會縮不會放**），字跟著變小，而下面留了
# 一大片空白。現場回報「字太小」。
#
# A3 橫式 1190.55 × 841.92 pt，四邊留 5mm：
#   可列印寬 ≈ 1162pt   可列印高 ≈ 813pt
# Excel 欄寬換算：pt = width × 7 × 0.75（見 _width_to_points）
# 列高同樣按天數分配，31 天與 28 天都要填滿整頁。
#
# 人數更多的單位自然會超出，那時才由 fitToPage 接手縮小。
# ⚠️ 邊界留 5mm 就好。openpyxl 預設左右各 0.75 吋（19mm），A3 橫式兩邊加起來
# 就吃掉 38mm，換算成欄寬是白白少掉一個多人的空間。5mm 是一般雷射印表機的
# 安全下限，再小會有印不到的風險。
MARGIN_INCH = 5.0 / 25.4

A3_W_PT = 1190.55
A3_H_PT = 841.92
_MARGIN_PT = MARGIN_INCH * 72
PRINTABLE_W_PT = A3_W_PT - 2 * _MARGIN_PT
PRINTABLE_H_PT = A3_H_PT - 2 * _MARGIN_PT

# ⚠️ **欄寬是按欄數分配的，不是固定值。**
#
# 派出所人員會調動，欄數每個月都可能不同。欄寬寫死的話：人多了會超出頁寬、
# 被 fitToPage 縮到看不清楚；人少了右邊留一大片空白。所以改成把可列印寬度
# 依權重分給各欄，人多自動變窄、人少自動變寬。
#
# 上下限是為了守住可讀性與美觀：
#   下限 3.8 → 約 20pt，兩位數代碼在 12pt 字下的最小可讀寬度
#   上限 13.0 → 人很少時不要讓格子胖到荒謬（手寫欄位寬一點無妨）
# 欄數多到連下限都排不下時（約 58 欄），才交給 fitToPage 整張縮。
MIN_COL_WIDTH = 3.8
MAX_COL_WIDTH = 13.0

# ⚠️ 欄寬權重在 lib/layout_model.column_weight，與 pdf_writer 共用——
# 兩邊用不同的權重，Excel 印出來就會跟 PDF 不一樣寬。

# ⚠️ **姓名列高度是算出來的，不能寫死。**
#
# 第一版寫死 62pt：三個字的姓名剛好，但「同仁專案臨檢」六個字直書就被切掉，
# 跨欄註記四行也只顯示得出兩行（「中班(17.18)」「限填1人」整個不見）。
# Excel 不會自動縮字，放不下就是切掉，而且**不會有任何警告**。
#
# 所以改成依實際內容算：取「最長的直書標題」與「註記行數」兩者所需高度的
# 較大者。
MIN_NAME_ROW_HEIGHT = 62
VERTICAL_LINE_RATIO = 1.35   # 直書一個字佔的高度 ÷ 字級（含字距）
NOTE_FONT_SIZE = 9           # 註記字小一級，四行才排得下
NOTE_LINE_RATIO = 1.45
CODE_ROW_HEIGHT = 20

# ⚠️ 沒有橫向標題列——標題是**最左邊那一整欄直書**（照紙本）。
ROW_NAME = 1
ROW_CODE = 2
ROW_FIRST_DAY = 3

_RED = "FFCC0000"
_BLUE = "FF1F4E9C"
_BLACK = "FF000000"

_THIN = Side(style="thin", color=_BLACK)
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
# shrink_to_fit：估算的字級塞不下時由 Excel 自己縮，保證不切字
_CENTER = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
# 姓名直書（Excel 的 textRotation 255 ＝ 直排）。
_VERTICAL = Alignment(horizontal="center", vertical="center", textRotation=255,
                      shrink_to_fit=True)
_NOTE = Alignment(horizontal="left", vertical="center", wrap_text=True)


# 換行格（不能配「縮小字型」）裡的半形字：Tahoma 粗體數字比半個字寬，照 _em_width 估
# 會被 Excel 折成「11／5」「2／7」（實測），多留四成五
DIGIT_EM_SLACK = 1.45
# 日期欄（縮小字型格）的半形寬度估計：比換行格的 1.45 緊，Excel 實測 1.2 時兩位數不會被縮
DATE_EM_SLACK = 1.2


def _digit_size(text: str, size: float, width_pt: float) -> float:
    return min(size, round(XLSX_FILL * width_pt / (_em_width(text) * DIGIT_EM_SLACK), 1))


def _em_width(text: str) -> float:
    """字串的寬度，以字級為單位：中文 1、半形 0.5。"""
    return sum(0.5 if ord(ch) < 128 else 1.0 for ch in text)


@dataclass(frozen=True)
class FontPlan:
    """一張表各類格子的字級（pt）與姓名列高度。Excel 版面由這裡決定。"""

    body: float        # 每天的格子（番號、休、空白欄）
    header: float      # 日期欄、星期欄的每天格子
    code: float        # 代碼列（21、A、早中晚）
    name: float        # 直書姓名（以人名欄寬為準）
    name_row: float    # 姓名列高度


def _fit(width_pt: float, height_pt: float, texts, digit_slack: float = 1.0) -> float:
    """讓 texts 裡最寬的字串橫書放進格子 XLSX_FILL 的字級。

    digit_slack：半形字串的寬度另外放大幾倍估（Tahoma 數字比半個字寬）。
    """
    widest = max((_em_width(t) * (digit_slack if t.isascii() else 1.0)
                  for t in texts if t), default=1.0)
    size = min(XLSX_FILL * width_pt / widest, XLSX_FILL * height_pt / LINE_RATIO)
    return round(min(MAX_FONT_SIZE, max(6.0, size)), 1)


def font_plan(sheet: Sheet) -> FontPlan:
    pairs = list(zip(sheet.columns, (_width_to_points(w) for w in column_widths(sheet))))

    def narrowest(pred):
        widths = [pt for column, pt in pairs if pred(column)]
        return min(widths) if widths else None

    member_pt = narrowest(lambda c: c.kind == COL_MEMBER)
    name = round(min(MAX_FONT_SIZE, XLSX_FILL * member_pt), 1) if member_pt else 12.0

    needed_header = 0.0
    for column in sheet.columns:
        if column.kind != COL_MEMBER or len(column.header) > NAME_ROW_CHARS:
            continue
        lines = _header_lines(column)
        if lines <= 1:
            continue
        need = lines * name * VERTICAL_LINE_RATIO
        if header_spans_code_row(column):
            need -= CODE_ROW_HEIGHT
        needed_header = max(needed_header, need)
    most_lines = max((len(b.note) for b in sheet.blocks), default=0)
    needed_note = most_lines * NOTE_FONT_SIZE * NOTE_LINE_RATIO
    name_row = max(MIN_NAME_ROW_HEIGHT, needed_header, needed_note)

    row_h = day_row_height(sheet.day_count, name_row)
    data = [c for c in sheet.columns if c.kind in (COL_MEMBER, COL_BLANK)]
    heads = [c for c in sheet.columns if c.kind not in (COL_MEMBER, COL_BLANK, COL_TITLE)]
    coded = [c for c in sheet.columns if c.code]
    body = _fit(narrowest(lambda c: c.kind in (COL_MEMBER, COL_BLANK)) or 30, row_h,
                {cell.text for c in data for cell in c.cells} | {"休", "00"})
    # ⚠️ 日期欄要照 Tahoma 實際字寬估：估太大時 Excel 的「縮小字型」只縮兩位數
    # （10～31），1～9 維持原字級，看起來大一截（維護者 2026-09-17 回報）。
    header = _fit(narrowest(lambda c: c in heads) or 30, row_h,
                  {cell.text for c in heads for cell in c.cells} | {"00"},
                  digit_slack=DATE_EM_SLACK)
    code = _fit(narrowest(lambda c: bool(c.code)) or 30, CODE_ROW_HEIGHT,
                {c.code for c in coded})
    return FontPlan(body=body, header=header, code=code, name=name, name_row=name_row)


def name_row_height(sheet: Sheet) -> float:
    """姓名列要多高才放得下直書姓名與最多行的註記（依 font_plan）。"""
    return font_plan(sheet).name_row


def _header_lines(column) -> int:
    """直書標題佔幾行：姓名每字一行，固定番／幹部的代碼另佔一行。"""
    extra = 1 if column.kind == COL_MEMBER and column.code else 0
    return len(column.header) + extra


def _vertical_size(column, width_pt: float, height_pt: float, preferred: float) -> float:
    """直書標題的字級：以寬度為主，字多到上下放不下時才縮（不切字）。"""
    by_height = height_pt / (_header_lines(column) * VERTICAL_LINE_RATIO)
    by_width = XLSX_FILL * width_pt
    return round(max(6.0, min(preferred, by_width, by_height)), 1)


# Excel 欄寬單位換算。
#
# ⚠️ **每欄不要再加 5px 的 padding。**
#
# 常見的公式寫成 `pixels = width × MDW + 5`，那是 Excel UI 顯示「8.43
# (64 像素)」時的算法。**實際版面佔的寬度是 `width × MDW`**——那 5px 在
# Excel 把「可見字元數」換算成儲存值的時候就已經算進去了。
#
# 第一版每欄多加 5px，48 欄就多算了 240px ≈ 64mm：程式以為排滿 410mm，
# Excel 實際只排了 346mm，預覽列印上左右各留了一大片白。現場回報「每格
# 寬度太小」就是這個。
#
# 是拿維護者的預覽列印截圖反推出來的：表格實際佔 346mm，而 48 欄的
# Σwidth × 7px = 1309px = 346mm，剛好吻合。
MAX_DIGIT_WIDTH_PX = 7      # Calibri 11 的最大數字寬（openpyxl 的預設字型）
PX_TO_PT = 0.75             # 96 dpi → 72 pt


def _width_to_points(width: float) -> float:
    return width * MAX_DIGIT_WIDTH_PX * PX_TO_PT


def _points_to_width(points: float) -> float:
    return points / (MAX_DIGIT_WIDTH_PX * PX_TO_PT)


def column_widths(sheet: Sheet) -> list[float]:
    """把可列印寬度依權重分給各欄，並夾在上下限之間。

    ⚠️ **夾到上下限的欄，差額要還給其他欄。**
    第一版夾完就算了，結果窄欄被夾寬時總寬會超出頁寬——40 人時就這樣多出
    3.6pt、整張表被 fitToPage 白白縮小一次。這裡改成反覆重分配：每輪把已經
    夾住的欄固定下來，剩下的寬度再按權重分給還沒夾住的欄，直到穩定。
    """
    weights = [column_weight(column) for column in sheet.columns]
    widths: list[float | None] = [None] * len(weights)

    remaining_pt = PRINTABLE_W_PT
    while True:
        free = [i for i, w in enumerate(widths) if w is None]
        if not free:
            break
        total_weight = sum(weights[i] for i in free)
        per_weight = remaining_pt / total_weight
        clamped = False
        for i in free:
            ideal = _points_to_width(per_weight * weights[i])
            bounded = min(MAX_COL_WIDTH, max(MIN_COL_WIDTH, ideal))
            if bounded != ideal:
                widths[i] = bounded
                remaining_pt -= _width_to_points(bounded)
                clamped = True
        if not clamped:
            for i in free:
                widths[i] = _points_to_width(per_weight * weights[i])
            break

    return [w for w in widths]


def day_row_height(day_count: int, name_height: float = MIN_NAME_ROW_HEIGHT) -> float:
    """日期列高：把剩下的高度分給每一天，讓 28 天的月份也填滿整頁。"""
    body = PRINTABLE_H_PT - name_height - CODE_ROW_HEIGHT
    return body / day_count


def fits_in_one_page(sheet: Sheet) -> bool:
    """欄數是否還排得下——False 表示要靠 fitToPage 整張縮小。"""
    return sum(_width_to_points(w) for w in column_widths(sheet)) <= PRINTABLE_W_PT + 1


def max_columns_per_page() -> int:
    """在最小欄寬下，一頁最多放得下幾欄。"""
    return int(PRINTABLE_W_PT // _width_to_points(MIN_COL_WIDTH))


def _border_range(ws: Worksheet, first_row, first_col, last_row, last_col) -> None:
    """替合併範圍內的**每一格**補上框線。

    ⚠️ Excel 不會自己補：框線只設在左上角那一格時，合併後只會畫出那一格的
    邊，其餘三邊是空的。標題欄（合併整欄）因此整欄看不到框線——現場回報
    「格線自己不見」。
    """
    for row in range(first_row, last_row + 1):
        for col in range(first_col, last_col + 1):
            ws.cell(row=row, column=col).border = _BORDER


def _argb(color: str) -> str:
    if color == RED:
        return _RED
    if color == BLUE:
        return _BLUE
    return _BLACK


def _setup_page(ws: Worksheet, sheet: Sheet) -> None:
    """A3 橫式、縮成一頁。

    ⚠️ 只設 ``page_setup`` 不夠——還要打開 ``sheetProperties.pageSetUpPr``
    的 fitToPage，否則 Excel 直接忽略縮放設定（PITFALLS XLS-2）。
    """
    ws.page_margins = PageMargins(
        left=MARGIN_INCH, right=MARGIN_INCH,
        top=MARGIN_INCH, bottom=MARGIN_INCH,
        header=0, footer=0,
    )
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = True

    # ⚠️ **放得下就固定 100%，不要交給 fitToPage。**
    #
    # 「調整成 1 頁寬 1 頁高」在算頁面分割時比實際保守：現場實測，表格自然
    # 尺寸 410 × 287mm、可列印區也是 410 × 287mm，關掉 fitToPage 用 100%
    # 印出來是紮紮實實的 1/1，但開著 fitToPage 它仍然縮了一級——上下貼滿、
    # **左右白掉一大片**。維護者回報「左右留的空間有點多」就是這個。
    #
    # 欄數真的超出容量時才讓它接手，那時縮小是應該的。
    if fits_in_one_page(sheet):
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_setup.fitToWidth = None
        ws.page_setup.fitToHeight = None
        ws.page_setup.scale = 100
    else:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1

    for index, width in enumerate(column_widths(sheet), start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _write_note(ws: Worksheet, first: int, block: Block) -> None:
    """區塊註記：跨該區塊所有欄的合併格，一行一筆純文字（一律黑字）。

    ⚠️ 註記不再帶顏色（維護者裁示 2026-09-16），所以也不必再走 RichText——
    openpyxl 一格只能一種字型，原本為了逐行上色才拆成 ``CellRichText``。
    """
    last = first + len(block.columns) - 1
    cell = ws.cell(row=ROW_NAME, column=first, value="\n".join(block.note))
    cell.font = Font(name=FONT_NAME, size=NOTE_FONT_SIZE)
    cell.alignment = _NOTE
    cell.border = _BORDER
    if last > first:
        _border_range(ws, ROW_NAME, first, ROW_NAME, last)
        ws.merge_cells(
            start_row=ROW_NAME, start_column=first,
            end_row=ROW_NAME, end_column=last,
        )


def title_sizes(pieces, width_pt: float) -> tuple[float, list[float]]:
    """回傳（中文字級, 每段字級）。純計算，測試直接驗「中文比數字大」「不超出欄高」。"""
    if not pieces:
        return 0.0, []
    size = round(min(TITLE_MAX_SIZE, XLSX_FILL * width_pt,
                     PRINTABLE_H_PT / (len(pieces) * VERTICAL_LINE_RATIO)), 1)
    return size, [_digit_size(p, size, width_pt) if p.isascii() else size for p in pieces]


def _write_title_column(ws: Worksheet, index: int, sheet: Sheet, width_pt: float) -> None:
    """最左邊那一整欄：直書標題，從姓名列一路合併到最後一天。"""
    last = ROW_FIRST_DAY - 1 + sheet.day_count
    # ⚠️ 不用 textRotation 直書：那會把「115」也拆成上下三個字。改成每段一行換行
    # 疊起來，數字那行照樣橫排（維護者 2026-09-17）。換行不能配「縮小字型」，
    # 字級要自己保證最寬的那段（115）塞得進欄寬。
    pieces = vertical_pieces(sheet.title)
    _, sizes = title_sizes(pieces, width_pt)
    cell = ws.cell(row=ROW_NAME, column=index)
    if pieces:
        # 每段換字型：數字段用 Tahoma。⚠️ 換行併在該段尾巴，單獨一段換行會讓檔案毀損
        end = len(pieces) - 1
        cell.value = CellRichText(*(
            TextBlock(InlineFont(rFont=_font_name(p), sz=sz, b=True),
                      p + ("" if i == end else "\n"))
            for i, (p, sz) in enumerate(zip(pieces, sizes))
        ))
    cell.alignment = _STACKED
    cell.border = _BORDER
    _border_range(ws, ROW_NAME, index, last, index)
    ws.merge_cells(
        start_row=ROW_NAME, start_column=index, end_row=last, end_column=index
    )


def _gapped_header(column, size: float) -> CellRichText:
    """「日　　期」：空白那段字級縮成 GAP_SCALE，間距才不會整整兩個字高。"""
    color = _argb(column.header_color)
    blocks = []
    for gap, chars in groupby(column.header, key=lambda char: char == GAP_CHAR):
        font = InlineFont(rFont=FONT_NAME, sz=round(size * GAP_SCALE, 1) if gap else size,
                          b=True, color=color)
        blocks.append(TextBlock(font, "".join(chars)))
    return CellRichText(*blocks)


# 固定番／幹部：姓名每字一行、代碼橫排一行，不用 textRotation（見 _write_name_with_code）
_STACKED = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_name_with_code(cell, column, size: float, width_pt: float) -> None:
    """姓名與代碼寫進同一格：姓名一字一行（看起來是直書），代碼橫排放上方或下方。

    ⚠️ 不能用 textRotation=255：Excel 一格只能一種文字方向，直排會把「21」也拆成
    上下兩個字。改用換行把中文字一個一個疊起來，代碼那一行照樣橫排（維護者 2026-09-16）。
    ⚠️ 換行（wrap_text）與「縮小字型以適合欄寬」不能同時用，所以字級要先算準
    （_vertical_size 以行數計）。姓名跟著女警變紅、代碼一律黑，用 RichText 分色。
    """
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    # ⚠️ 換行要併進前後文字：只含換行的一段存檔時被當空白吃掉，Excel 判定檔案毀損打不開
    name_font = InlineFont(rFont=FONT_NAME, sz=size, b=True, color=_argb(column.header_color))
    code_size = _digit_size(column.code, size, width_pt)
    code_font = InlineFont(rFont=_font_name(column.code), sz=code_size, b=True, color=_BLACK)
    name = "\n".join(column.header)
    if column.code_above:
        # 換行歸姓名那段：換行跟著代號字型的話，代號和名字之間會多空一截
        # （維護者 2026-09-17：「貌合神離」）
        parts = (TextBlock(code_font, column.code), TextBlock(name_font, "\n" + name))
    else:
        parts = (TextBlock(name_font, name + "\n"), TextBlock(code_font, column.code))
    cell.value = CellRichText(*parts)
    cell.alignment = _STACKED


def _write_column(
    ws: Worksheet, index: int, column, day_count: int, plan: FontPlan,
    width_pt: float, skip_name: bool = False,
) -> None:
    """``skip_name`` 為真時不畫姓名列——那一格被區塊註記的合併格佔走了。"""
    cell_size = plan.body if column.kind in (COL_MEMBER, COL_BLANK) else plan.header
    if not skip_name:
        spans = header_spans_code_row(column)
        header_h = plan.name_row + (CODE_ROW_HEIGHT if spans else 0)
        with_code = column.kind == COL_MEMBER and bool(column.code)
        if len(column.header) > 1 or with_code:
            preferred = plan.name if column.kind == COL_MEMBER else MAX_FONT_SIZE
            size = _vertical_size(column, width_pt, header_h, preferred)
        else:
            size = _fit(width_pt, header_h, {column.header})
        header = ws.cell(row=ROW_NAME, column=index)
        if with_code:
            _write_name_with_code(header, column, size, width_pt)
        else:
            header.value = column.header or None
            header.font = Font(
                name=FONT_NAME, size=size,
                color=_argb(column.header_color), bold=True,
            )
            if GAP_CHAR in column.header and len(column.header) > 1:
                header.value = _gapped_header(column, size)
            # ⚠️ 欄很窄，多字標題橫著放會被切掉（「快打勤務」「日期」都踩過）
            # ——只要超過一個字就直書。
            header.alignment = _VERTICAL if len(column.header) > 1 else _CENTER
        header.border = _BORDER

        # 代碼列沒東西的欄，標題跨姓名列與代碼列合併（規則與 PDF 共用）
        if column.kind == COL_TITLE or header_spans_code_row(column):
            _border_range(ws, ROW_NAME, index, ROW_CODE, index)
            ws.merge_cells(
                start_row=ROW_NAME, start_column=index,
                end_row=ROW_CODE, end_column=index,
            )
        else:
            # ⚠️ 只有姓名跟著女警變紅，番號一律黑的。
            code = ws.cell(row=ROW_CODE, column=index, value=column.code or None)
            code.font = Font(name=_font_name(column.code), size=plan.code)
            code.alignment = _CENTER
            code.border = _BORDER
    else:
        code = ws.cell(row=ROW_CODE, column=index, value=column.code or None)
        code.font = Font(name=_font_name(column.code), size=plan.code, color=_RED)
        code.alignment = _CENTER
        code.border = _BORDER

    for day in range(day_count):
        model = column.cells[day]
        cell = ws.cell(row=ROW_FIRST_DAY + day, column=index, value=model.text or None)
        cell.font = Font(name=_font_name(model.text), size=cell_size, color=_argb(model.color))
        cell.alignment = _CENTER
        cell.border = _BORDER


def write_sheet(sheet: Sheet, path: str) -> None:
    """把版面模型寫成 xlsx。"""
    wb = Workbook()
    ws = wb.active
    ws.title = f"{sheet.month}月"

    columns = sheet.columns
    _setup_page(ws, sheet)

    plan = font_plan(sheet)
    widths_pt = [_width_to_points(w) for w in column_widths(sheet)]
    name_height = plan.name_row
    ws.row_dimensions[ROW_NAME].height = name_height
    ws.row_dimensions[ROW_CODE].height = CODE_ROW_HEIGHT
    row_height = day_row_height(sheet.day_count, name_height)
    for day in range(sheet.day_count):
        ws.row_dimensions[ROW_FIRST_DAY + day].height = row_height

    index = 1
    for block in sheet.blocks:
        if block.note:
            _write_note(ws, index, block)
        for column in block.columns:
            if column.kind == COL_TITLE:
                _write_title_column(ws, index, sheet, widths_pt[index - 1])
            else:
                _write_column(
                    ws, index, column, sheet.day_count, plan, widths_pt[index - 1],
                    skip_name=bool(block.note),
                )
            index += 1

    # ⚠️ 不設凍結窗格（維護者裁示）。這是一張要列印的表，不是拿來捲動看的；
    # 凍結線在畫面上多一條橫槓，反而干擾。
    wb.save(path)


def member_column_count(sheet: Sheet) -> int:
    """整張表有幾欄是人（測試與版面估算用）。"""
    return sum(1 for column in sheet.columns if column.kind == COL_MEMBER)
